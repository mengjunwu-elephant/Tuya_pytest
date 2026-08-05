# -*- coding: utf-8 -*-
"""Tuya upper_body 接口响应时间诊断脚本。

参考 elephant-pytest/tools/diagnostics/response_time.py：
对参数表中的每个接口循环调用，统计最大/最小/平均/中位数/方差/标准差/错误率/丢包率，
并写入 RK3562 用例表的「python接口响应时间」sheet。
"""
from __future__ import annotations

import argparse
import json
import statistics
import sys
import time
from pathlib import Path
from typing import Any, Callable

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook, load_workbook
from pytuyarobot.command_result import CommandResult

from settings import TuyaRobotBase, TuyaConnectionConfig

TOOLS_DIR = Path(__file__).resolve().parent
DEFAULT_PARAM_FILE = TOOLS_DIR / "upper_body_response_time.xlsx"
REPORT_SHEET_NAME = "python接口响应时间"
ZERO_ANGLES = list(TuyaRobotBase.UPPER_BODY_ZERO_ANGLES)
SPEED = int(TuyaRobotBase.speed)

# 占位符：脚本在 go_zero 后用实机回读坐标替换
LEFT_COORDS_TOKEN = "$LEFT_COORDS$"
RIGHT_COORDS_TOKEN = "$RIGHT_COORDS$"

# (api, target, args_json_obj, note)
API_ROWS: list[tuple[str, str, Any, str]] = [
    ("get_upper_main_version", "both", [], "获取主控版本"),
    ("get_upper_modify_version", "both", [], "获取修改版本号"),
    ("get_upper_debug_state", "both", [], "读取debug日志模式"),
    ("set_upper_debug_state", "both", [0], "设置debug日志模式"),
    ("get_upper_fresh_mode", "both", [], "读取运动模式"),
    ("set_upper_fresh_mode", "both", [0], "设置运动模式为插补0"),
    ("get_upper_angles", "both", [], "获取全关节角度"),
    ("get_upper_coords", "both", [], "获取笛卡尔坐标"),
    ("get_upper_robot_status", "both", [], "读取机器人状态"),
    ("get_upper_is_init_calibrate", "both", [], "读取零位校准状态"),
    ("get_upper_zero_encoder", "both", [], "读取零位编码器值"),
    ("get_upper_plan_sp", "both", [0], "读取规划速度mode0"),
    ("get_upper_plan_acc", "both", [0], "读取规划加速度mode0"),
    ("get_upper_joints_status", "both", [], "读取关节状态"),
    ("get_upper_joints_current", "both", [], "读取关节电流"),
    ("get_upper_joints_run_sp", "both", [], "读取关节运行速度"),
    ("get_upper_encoders", "both", [], "读取编码器"),
    ("get_upper_joint_loss_count", "both", [1], "读取J1通讯丢包计数"),
    ("get_upper_model_direction", "both", [], "读取模型方向"),
    ("get_upper_tool_reference", "both", [], "读取工具坐标系"),
    ("get_upper_reference_frame", "both", [], "读取参考坐标系"),
    ("get_upper_movement_type", "both", [], "读取运动类型"),
    ("get_upper_end_type", "both", [], "读取末端类型"),
    ("get_upper_vr_mode", "both", [], "读取VR模式"),
    ("get_coords", "both", [], "读取笛卡尔坐标get_coords"),
    ("get_upper_collision_mode", "both", [], "读取碰撞模式"),
    ("get_upper_collision_threshold", "both", [], "读取碰撞阈值"),
    ("set_upper_plan_sp", "both", [0, SPEED], "设置规划速度mode0"),
    ("set_upper_plan_acc", "both", [0, SPEED], "设置规划加速度mode0"),
    ("upper_pause", "both", [], "暂停运动"),
    ("upper_resume", "both", [], "恢复运动"),
    ("set_upper_joint_enable", "both", [1, 1], "设置J1使能"),
    ("upper_set_break", "both", [1, 0], "设置J1抱闸松开"),
    (
        "send_upper_angle",
        "left",
        {"args": [1, 0.0, SPEED], "kwargs": {"_async": True}},
        "左臂J1零位发令_只测RTT",
    ),
    (
        "send_upper_angles",
        "both",
        {
            "args": [ZERO_ANGLES, SPEED, SPEED, ZERO_ANGLES, SPEED, SPEED],
            "kwargs": {"_async": True},
        },
        "双臂零位发令_只测RTT",
    ),
    (
        "send_upper_coords",
        "both",
        {
            "args": [LEFT_COORDS_TOKEN, SPEED, RIGHT_COORDS_TOKEN, SPEED],
            "kwargs": {"_async": True},
        },
        "双臂当前坐标重发_只测RTT",
    ),
    ("get_upper_is_in_position", "both", [0], "查询是否到位mode0"),
    ("get_upper_is_moving", "both", [], "查询是否运动中"),
    ("is_upper_powered_on", "both", [], "查询上电状态"),
    ("clear_upper_error", "both", [254], "清除错误"),
    (
        "upper_go_zero",
        "both",
        {"args": [], "kwargs": {"_async": True}},
        "执行回零发令_只测RTT",
    ),
    ("get_upper_is_paused", "both", [], "查询是否暂停"),
    ("upper_stop", "both", [], "停止运动"),
    ("get_upper_joint_acc", "both", [], "读取关节加速度"),
    ("get_upper_joints_max_angle", "both", [], "读取关节最大角度"),
    ("get_upper_joints_min_angle", "both", [], "读取关节最小角度"),
    ("set_upper_joint_acc", "both", [1, 50.0], "设置J1加速度"),
    (
        "set_joint_max_angle",
        "both",
        [1, TuyaRobotBase.UPPER_BODY_JOINT_SOFT_LIMITS[1][1]],
        "重设J1软件上限",
    ),
    (
        "set_joint_min_angle",
        "both",
        [1, TuyaRobotBase.UPPER_BODY_JOINT_SOFT_LIMITS[1][0]],
        "重设J1软件下限",
    ),
    ("set_upper_vr_mode", "both", [0, 0], "设置VR模式关闭"),
    ("get_upper_filter_len", "both", [1], "读取滤波长度rank1"),
    ("set_upper_filter_len", "both", [1, 0], "设置滤波长度rank1"),
    ("set_upper_movement_type", "both", [0, 0], "设置运动类型"),
    ("set_upper_end_type", "both", [0, 0], "设置末端类型"),
]


def is_packet_loss(result: Any) -> bool:
    if result is None:
        return True
    if result == "":
        return True
    if result == () or result == []:
        return True
    return False


def is_error_result(result: Any) -> bool:
    if isinstance(result, CommandResult) and not result.ok:
        return True
    return False


def measure_time(func: Callable[[], Any], times: int = 1000) -> dict[str, Any]:
    """循环调用并统计响应时间（毫秒）。"""
    packet_losses = 0
    error_times = 0
    valid_times: list[float] = []

    for i in range(times):
        start = time.perf_counter()
        try:
            result = func()
        except Exception as exc:  # noqa: BLE001 — 诊断脚本需统计任意异常为错误
            end = time.perf_counter()
            res_time = round((end - start) * 1000, 3)
            error_times += 1
            print(f"****** 第{i}次异常 {res_time} ms: {exc!r} ******")
            continue
        end = time.perf_counter()
        res_time = round((end - start) * 1000, 3)
        print(f"****** 第{i}次运行 {res_time} ms, 结果={result!r} ******")

        if is_packet_loss(result):
            packet_losses += 1
        elif is_error_result(result):
            error_times += 1
        else:
            valid_times.append(res_time)

    valid_count = len(valid_times)
    if valid_count > 0:
        average_time = round(sum(valid_times) / valid_count, 3)
        min_time = min(valid_times)
        max_time = max(valid_times)
        variance = round(statistics.variance(valid_times), 3) if valid_count > 1 else 0
        std_dev = round(statistics.stdev(valid_times), 3) if valid_count > 1 else 0
        median_time = round(statistics.median(valid_times), 3)
    else:
        average_time = min_time = max_time = variance = std_dev = median_time = None

    return {
        "times": times,
        "valid_times": valid_count,
        "average": average_time,
        "min": min_time,
        "max": max_time,
        "median": median_time,
        "variance": variance,
        "std_dev": std_dev,
        "packet_loss_rate": round((packet_losses / times) * 100, 3) if times else 0,
        "error_rate": round((error_times / times) * 100, 3) if times else 0,
    }


def parse_args_json(raw: Any) -> tuple[list[Any], dict[str, Any]]:
    if raw is None or raw == "":
        return [], {}
    if isinstance(raw, (list, tuple)):
        return list(raw), {}
    if isinstance(raw, dict):
        data = raw
    else:
        data = json.loads(str(raw))
    if isinstance(data, list):
        return data, {}
    if not isinstance(data, dict):
        raise ValueError(f"args_json 必须是 list 或 dict: {raw!r}")
    args = list(data.get("args", []))
    kwargs = dict(data.get("kwargs", {}))
    # 兼容直接写关键字且无 args 键
    if "args" not in data and "kwargs" not in data:
        kwargs = dict(data)
        args = []
    return args, kwargs


def resolve_tokens(value: Any, left_coords: list[float] | None, right_coords: list[float] | None) -> Any:
    if isinstance(value, str):
        if value == LEFT_COORDS_TOKEN:
            if left_coords is None:
                raise RuntimeError("未准备左臂坐标，无法替换 $LEFT_COORDS$")
            return list(left_coords)
        if value == RIGHT_COORDS_TOKEN:
            if right_coords is None:
                raise RuntimeError("未准备右臂坐标，无法替换 $RIGHT_COORDS$")
            return list(right_coords)
        return value
    if isinstance(value, list):
        return [resolve_tokens(v, left_coords, right_coords) for v in value]
    if isinstance(value, dict):
        return {k: resolve_tokens(v, left_coords, right_coords) for k, v in value.items()}
    return value


def resolve_target(device: TuyaRobotBase, target: str):
    key = (target or "both").strip().lower()
    mapping = {
        "left": device.left_arm,
        "right": device.right_arm,
        "both": device.upper_body,
        "upper": device.upper_body,
        "upper_body": device.upper_body,
        "robot": device.robot,
    }
    if key not in mapping:
        raise ValueError(f"不支持的 target: {target!r}")
    return mapping[key]


def read_param_rows(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, data_only=True)
    if "apis" not in wb.sheetnames:
        raise RuntimeError(f"参数表缺少 sheet 'apis': {path}")
    ws = wb["apis"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    header_map = {str(h).strip(): idx for idx, h in enumerate(headers) if h}
    required = ("api", "enabled", "target", "times", "args_json")
    for name in required:
        if name not in header_map:
            raise RuntimeError(f"参数表缺少列 {name}: {path}")

    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        api = row[header_map["api"]]
        if api is None or str(api).strip() == "":
            continue
        enabled_raw = row[header_map["enabled"]]
        enabled = str(enabled_raw).strip() in {"1", "1.0", "true", "True", "YES", "yes"}
        times_raw = row[header_map["times"]]
        times = int(times_raw) if times_raw not in (None, "") else 1000
        note = ""
        if "note" in header_map and row[header_map["note"]] is not None:
            note = str(row[header_map["note"]])
        rows.append(
            {
                "api": str(api).strip(),
                "enabled": enabled,
                "target": str(row[header_map["target"]] or "both").strip(),
                "times": times,
                "args_json": row[header_map["args_json"]],
                "note": note,
            }
        )
    wb.close()
    return rows


def init_param_excel(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "apis"
    ws.append(["api", "enabled", "target", "times", "args_json", "note"])
    for api, target, args_obj, note in API_ROWS:
        ws.append([api, 1, target, 1000, json.dumps(args_obj, ensure_ascii=False), note])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"已生成参数表: {path} ({len(API_ROWS)} 行)")


def find_report_workbook(explicit: str | None = None) -> Path:
    if explicit:
        path = Path(explicit)
        if not path.exists():
            raise FileNotFoundError(path)
        return path
    matches = sorted(Path(r"C:\Users\HP\Desktop").glob("**/Tuya_Rk3562_v2.0*.xlsx"))
    # 排除临时/锁文件
    matches = [m for m in matches if not m.name.startswith("~$")]
    if not matches:
        raise FileNotFoundError("未找到 Tuya_Rk3562_v2.0*.xlsx，请用 --report 指定路径")
    return matches[0]


def ensure_report_headers(ws) -> dict[str, int]:
    """确保表头含标准差列，返回列名 -> 1-based 列号。"""
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    # 去掉尾部空列
    while headers and (headers[-1] is None or str(headers[-1]).strip() == ""):
        headers.pop()

    expected = [
        "接口名称",
        "测试接口",
        "python运行1000次平均响应时间（ms）",
        "python运行1000次最大最小值（ms）",
        "python运行1000次中位数值（ms）",
        "python运行1000次方差",
        "python运行1000次标准差",
        "python丢包率",
        "python错误率",
        "python测试截图",
    ]

    # 若缺少标准差列，插入到方差之后
    names = [str(h).strip() if h is not None else "" for h in headers]
    if "python运行1000次标准差" not in names:
        insert_at = None
        for idx, name in enumerate(names, start=1):
            if "方差" in name:
                insert_at = idx + 1
                break
        if insert_at is None:
            insert_at = len(names) + 1
        ws.insert_cols(insert_at)
        ws.cell(1, insert_at).value = "python运行1000次标准差"
        names = [str(ws.cell(1, c).value or "").strip() for c in range(1, max(len(names) + 1, insert_at) + 1)]
        # 重新读取
        names = []
        col = 1
        while True:
            val = ws.cell(1, col).value
            if val is None and col > 10:
                break
            if val is not None:
                names.append(str(val).strip())
            col += 1
            if col > 20:
                break

    # 若整表几乎为空，写入完整表头
    if not any(names):
        for idx, title in enumerate(expected, start=1):
            ws.cell(1, idx).value = title
        names = expected[:]

    # 仍缺则补全缺失列到末尾
    existing = {str(ws.cell(1, c).value).strip(): c for c in range(1, 21) if ws.cell(1, c).value}
    next_col = max(existing.values(), default=0) + 1
    for title in expected:
        if title not in existing:
            ws.cell(1, next_col).value = title
            existing[title] = next_col
            next_col += 1

    # 模糊匹配关键列
    col_map: dict[str, int] = {}
    for c in range(1, 21):
        val = ws.cell(1, c).value
        if val is None:
            continue
        text = str(val).strip()
        col_map[text] = c
        if text == "接口名称":
            col_map["name"] = c
        elif text == "测试接口":
            col_map["api"] = c
        elif "平均" in text:
            col_map["average"] = c
        elif "最大最小" in text:
            col_map["minmax"] = c
        elif "中位" in text:
            col_map["median"] = c
        elif "标准差" in text:
            col_map["std_dev"] = c
        elif "方差" in text:
            col_map["variance"] = c
        elif "丢包" in text:
            col_map["packet_loss"] = c
        elif "错误" in text:
            col_map["error"] = c
        elif "截图" in text:
            col_map["screenshot"] = c
    return col_map


def write_report_rows(
    report_path: Path,
    results: list[dict[str, Any]],
) -> None:
    wb = load_workbook(report_path)
    if REPORT_SHEET_NAME not in wb.sheetnames:
        raise RuntimeError(f"报告文件缺少 sheet {REPORT_SHEET_NAME!r}: {report_path}")
    ws = wb[REPORT_SHEET_NAME]
    col_map = ensure_report_headers(ws)

    for idx, item in enumerate(results):
        row = idx + 2
        stats = item["stats"]
        if "name" in col_map:
            ws.cell(row, col_map["name"]).value = item.get("note") or item["api"]
        if "api" in col_map:
            ws.cell(row, col_map["api"]).value = item["api"]
        if "average" in col_map:
            ws.cell(row, col_map["average"]).value = stats["average"]
        if "minmax" in col_map:
            if stats["max"] is None or stats["min"] is None:
                ws.cell(row, col_map["minmax"]).value = None
            else:
                ws.cell(row, col_map["minmax"]).value = f"{stats['max']} / {stats['min']}"
        if "median" in col_map:
            ws.cell(row, col_map["median"]).value = stats["median"]
        if "variance" in col_map:
            ws.cell(row, col_map["variance"]).value = stats["variance"]
        if "std_dev" in col_map:
            ws.cell(row, col_map["std_dev"]).value = stats["std_dev"]
        if "packet_loss" in col_map:
            ws.cell(row, col_map["packet_loss"]).value = f"{stats['packet_loss_rate']}%"
        if "error" in col_map:
            ws.cell(row, col_map["error"]).value = f"{stats['error_rate']}%"

    wb.save(report_path)
    print(f"已写入报告: {report_path} / {REPORT_SHEET_NAME} ({len(results)} 行)")


def snapshot_coords(device: TuyaRobotBase) -> tuple[list[float], list[float]]:
    raw = device.result_data(device.upper_body.get_upper_coords())
    if not isinstance(raw, dict) or "left" not in raw or "right" not in raw:
        raise RuntimeError(f"get_upper_coords 返回格式错误: {raw!r}")
    left = [float(x) for x in raw["left"]]
    right = [float(x) for x in raw["right"]]
    if len(left) != 6 or len(right) != 6:
        raise RuntimeError(f"坐标长度错误: left={left!r}, right={right!r}")
    return left, right


def print_stats(api: str, stats: dict[str, Any]) -> None:
    print("\n========= 统计结果 =========")
    print(f"接口: {api}")
    print(f"总运行次数: {stats['times']}")
    print(f"有效次数: {stats['valid_times']}")
    print(f"平均响应时间: {stats['average']} ms")
    print(f"最大值: {stats['max']} ms")
    print(f"最小值: {stats['min']} ms")
    print(f"中位数: {stats['median']} ms")
    print(f"方差: {stats['variance']}")
    print(f"标准差: {stats['std_dev']}")
    print(f"丢包率: {stats['packet_loss_rate']} %")
    print(f"错误率: {stats['error_rate']} %")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Tuya upper_body 响应时间诊断")
    parser.add_argument(
        "--param",
        default=str(DEFAULT_PARAM_FILE),
        help="参数 Excel 路径",
    )
    parser.add_argument(
        "--report",
        default=None,
        help="结果 Excel 路径（默认自动查找 Tuya_Rk3562_v2.0*.xlsx）",
    )
    parser.add_argument("--times", type=int, default=None, help="覆盖参数表次数")
    parser.add_argument(
        "--apis",
        default=None,
        help="仅测指定接口，逗号分隔",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="只校验参数表，不连接设备、不写报告",
    )
    parser.add_argument(
        "--init-excel",
        action="store_true",
        help="按白名单生成/覆盖参数表后退出",
    )
    parser.add_argument(
        "--skip-go-zero",
        action="store_true",
        help="跳过启动时 go_zero",
    )
    parser.add_argument(
        "--no-write-report",
        action="store_true",
        help="只打印统计，不写报告 Excel",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    param_path = Path(args.param)

    if args.init_excel:
        init_param_excel(param_path)
        return 0

    if not param_path.exists():
        print(f"参数表不存在，先生成: {param_path}")
        init_param_excel(param_path)

    rows = read_param_rows(param_path)
    api_filter = None
    if args.apis:
        api_filter = {name.strip() for name in args.apis.split(",") if name.strip()}

    selected = []
    for row in rows:
        if not row["enabled"]:
            continue
        if api_filter and row["api"] not in api_filter:
            continue
        selected.append(row)

    print(f"参数表: {param_path}")
    print(f"待测接口数: {len(selected)}")
    for row in selected:
        print(f"  - {row['api']} target={row['target']} times={args.times or row['times']}")

    if args.dry_run:
        # 校验 args_json 可解析
        for row in selected:
            parse_args_json(row["args_json"])
        print("dry-run 通过：参数表结构与 args_json 均可解析")
        return 0

    report_path = None
    if not args.no_write_report:
        report_path = find_report_workbook(args.report)
        print(f"报告文件: {report_path}")

    device = TuyaRobotBase(TuyaConnectionConfig.from_env())
    left_coords = right_coords = None
    results: list[dict[str, Any]] = []
    try:
        if not args.skip_go_zero:
            print("执行 go_zero 到零位…")
            device.go_zero()
        print("回读零位坐标供 send_upper_coords 使用…")
        left_coords, right_coords = snapshot_coords(device)
        print(f"left_coords={left_coords}")
        print(f"right_coords={right_coords}")

        # 运动发令使用异步，避免同步等待到位
        try:
            device.upper_body.set_upper_motion_async(True)
        except Exception as exc:  # noqa: BLE001
            print(f"警告: set_upper_motion_async(True) 失败: {exc!r}")

        for row in selected:
            api = row["api"]
            times = args.times if args.times is not None else row["times"]
            target_obj = resolve_target(device, row["target"])
            if not hasattr(target_obj, api):
                print(f"跳过 {api}: 目标对象 {row['target']} 无此方法")
                continue
            method = getattr(target_obj, api)
            raw_args, raw_kwargs = parse_args_json(row["args_json"])
            call_args = resolve_tokens(raw_args, left_coords, right_coords)
            call_kwargs = resolve_tokens(raw_kwargs, left_coords, right_coords)

            print(f"\n>>> 开始测试 {api} x {times}, target={row['target']}")
            print(f"args={call_args!r} kwargs={call_kwargs!r}")

            def _call(
                _method=method,
                _args=call_args,
                _kwargs=call_kwargs,
            ):
                return _method(*_args, **_kwargs)

            stats = measure_time(_call, times=times)
            print_stats(api, stats)
            results.append({"api": api, "note": row.get("note") or api, "stats": stats})
    finally:
        try:
            device.upper_body.set_upper_motion_async(False)
        except Exception:  # noqa: BLE001
            pass
        device.close()
        print("已关闭设备连接")

    if report_path is not None and results:
        write_report_rows(report_path, results)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
