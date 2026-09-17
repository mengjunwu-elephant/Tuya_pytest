"""线程安全的统计、明细和汇总报告。"""

from __future__ import annotations

import os
import threading
import time
import traceback
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

from openpyxl import Workbook

from ..utils import excel_value, numeric_values, utc_text


class ReportCollector:
    SHEET_HEADERS = {
        "上半身运动": (
            "时间", "统一循环", "动作编号", "接口", "目标", "参数",
            "期望", "实际", "偏差", "耗时秒", "是否到位",
            "是否越界", "结果", "错误",
        ),
        "头部运动": (
            "时间", "统一循环", "动作编号", "接口", "目标", "参数",
            "期望", "实际", "偏差", "耗时秒", "是否到位",
            "是否越界", "结果", "错误",
        ),
        "底盘运动": (
            "时间", "统一循环", "动作编号", "接口", "前进速度mps",
            "旋转速度rads", "持续秒", "返回", "结果", "错误",
        ),
        "上半身遥测": (
            "时间", "角度", "坐标", "关节电流", "运行速度", "编码器",
            "丢包计数", "关节状态", "机器人状态", "运动状态",
            "暂停状态", "错误",
        ),
        "头部遥测": (
            "时间", "上电状态", "角度", "运动状态", "错误",
        ),
        "底盘遥测": (
            "时间", "上电状态", "错误",
        ),
        "异常事件": (
            "时间", "级别", "子系统", "阶段", "信息", "异常", "堆栈",
        ),
    }

    def __init__(self, root: Path, config: dict[str, Any]) -> None:
        self.root = root
        self.root.mkdir(parents=True, exist_ok=True)
        self.config = config
        self.lock = threading.RLock()
        self.stats: dict[str, dict[str, float]] = defaultdict(
            lambda: defaultdict(float)
        )
        self.started_monotonic = time.monotonic()
        self.started_text = utc_text()
        self._hour_key = ""
        self._detail_path: Path | None = None
        self._wb: Workbook | None = None
        self._sheets: dict[str, Any] = {}
        self._new_detail_workbook_locked()

    def _new_detail_workbook_locked(self) -> None:
        hour_key = datetime.now().strftime("%Y%m%d_%H")
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheets = {}
        for name, headers in self.SHEET_HEADERS.items():
            sheet = workbook.create_sheet(name)
            sheet.append(list(headers))
            sheet.freeze_panes = "A2"
            sheets[name] = sheet
        self._hour_key = hour_key
        self._detail_path = self.root / f"details_{hour_key}.xlsx"
        self._wb = workbook
        self._sheets = sheets

    def append(self, sheet_name: str, row: Sequence[Any]) -> None:
        with self.lock:
            current_hour = datetime.now().strftime("%Y%m%d_%H")
            if current_hour != self._hour_key:
                self._save_detail_locked()
                self._new_detail_workbook_locked()
            expected = len(self.SHEET_HEADERS[sheet_name])
            if len(row) != expected:
                raise ValueError(
                    f"{sheet_name} 报告列数错误，期望 {expected}，实际 {len(row)}"
                )
            self._sheets[sheet_name].append(
                [excel_value(value) for value in row]
            )

    def count(
        self, category: str, metric: str, amount: float = 1.0
    ) -> None:
        with self.lock:
            self.stats[category][metric] += amount

    def set_max(self, category: str, metric: str, value: float) -> None:
        with self.lock:
            bucket = self.stats[category]
            bucket[metric] = max(bucket.get(metric, value), value)

    def observe_numbers(
        self, category: str, metric: str, value: Any
    ) -> None:
        values = numeric_values(value)
        if not values:
            return
        with self.lock:
            bucket = self.stats[category]
            bucket[f"{metric}样本数"] += len(values)
            bucket[f"{metric}累计值"] += sum(values)
            minimum_key = f"{metric}最小值"
            maximum_key = f"{metric}最大值"
            bucket[minimum_key] = min(
                bucket.get(minimum_key, min(values)), min(values)
            )
            bucket[maximum_key] = max(
                bucket.get(maximum_key, max(values)), max(values)
            )

    def event(
        self,
        level: str,
        subsystem: str,
        phase: str,
        message: str,
        exc: BaseException | None = None,
    ) -> None:
        stack = ""
        if exc is not None:
            stack = "".join(
                traceback.format_exception(type(exc), exc, exc.__traceback__)
            )
        self.append(
            "异常事件",
            (
                utc_text(), level, subsystem, phase, message,
                repr(exc) if exc else "", stack,
            ),
        )
        self.count("事件", level)

    def _save_detail_locked(self) -> None:
        if self._wb is None or self._detail_path is None:
            return
        temp = self._detail_path.with_name(
            f".{self._detail_path.stem}.tmp.xlsx"
        )
        self._wb.save(temp)
        os.replace(temp, self._detail_path)

    def _save_summary_locked(self) -> None:
        workbook = Workbook()
        summary = workbook.active
        summary.title = "汇总"
        summary.append(("项目", "指标", "数值"))
        summary.append(("运行", "开始时间", self.started_text))
        summary.append(
            (
                "运行",
                "累计运行秒",
                round(time.monotonic() - self.started_monotonic, 3),
            )
        )
        for category in sorted(self.stats):
            bucket = self.stats[category]
            for metric in sorted(bucket):
                summary.append((category, metric, bucket[metric]))
        config_sheet = workbook.create_sheet("配置")
        config_sheet.append(("键", "值"))
        for key in sorted(self.config):
            config_sheet.append((key, excel_value(self.config[key])))
        path = self.root / "summary.xlsx"
        temp = path.with_name(".summary.tmp.xlsx")
        workbook.save(temp)
        os.replace(temp, path)

    def save(self) -> None:
        with self.lock:
            self._save_detail_locked()
            self._save_summary_locked()
