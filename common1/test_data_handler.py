# -*- coding: utf-8 -*-
from __future__ import annotations

import os
from typing import Any, Optional, Sequence

from openpyxl import load_workbook


def get_test_data_from_excel(
    file: str,
    sheet_name: str,
    required_columns: Optional[Sequence[str]] = None,
) -> list[dict[str, Any]]:
    """读取首行为字段名的 Excel Sheet，并返回用例字典列表。

    空行会被跳过；尾部仅因格式产生的空列会被安全忽略。

    :param file: xlsx 文件路径。
    :param sheet_name: Sheet 名称，应与对应 SDK 接口名一致。
    :param required_columns: 可选的必填列名集合。
    :raises FileNotFoundError: 文件不存在。
    :raises KeyError: Sheet 不存在。
    :raises ValueError: 表头为空、包含空列名或缺少必填列。
    """
    if not os.path.isfile(file):
        raise FileNotFoundError(file)

    workbook = load_workbook(file, read_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise KeyError(
                f"工作表 {sheet_name!r} 不存在，当前工作簿包含：{workbook.sheetnames!r}"
            )

        sheet = workbook[sheet_name]
        keys = [sheet.cell(1, index).value for index in range(1, sheet.max_column + 1)]
        while keys and (
            keys[-1] is None
            or (isinstance(keys[-1], str) and keys[-1].strip() == "")
        ):
            keys.pop()

        if not keys:
            raise ValueError("Excel 首行没有任何列名")
        if any(
            key is None or (isinstance(key, str) and key.strip() == "")
            for key in keys
        ):
            raise ValueError("Excel 首行存在空列名，请删除空列或填写表头")

        column_names = [str(key).strip() for key in keys]
        if required_columns is not None:
            missing = set(required_columns) - set(column_names)
            if missing:
                raise ValueError(f"Excel 缺少必填列：{sorted(missing)}")

        cases: list[dict[str, Any]] = []
        for row_index in range(2, sheet.max_row + 1):
            case = {
                column_name: sheet.cell(row_index, column_index).value
                for column_index, column_name in enumerate(column_names, 1)
            }
            if all(
                value is None
                or (isinstance(value, str) and value.strip() == "")
                for value in case.values()
            ):
                continue
            cases.append(case)
        return cases
    finally:
        workbook.close()
